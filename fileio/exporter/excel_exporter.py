from tkinter import filedialog

import openpyxl
from openpyxl.utils import get_column_letter
from tkinter import messagebox

class ExcelExporter:
    def __init__(self, data, columns):
        """
        :param data: [{col:value, ...}, ...] 형태의 리스트
        :param columns: 컬럼명 리스트
        """
        self.data = data
        self.columns = columns

    def export(self):
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel 파일", "*.xlsx")],
            title="엑셀 파일로 저장"
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Pole Summary"

            # 컬럼 제목
            ws.append(self.columns)

            # 데이터 추가
            for row in self.data:
                ws.append([row.get(col, "") for col in self.columns])

            # 컬럼 너비 자동 조정
            for col_num, col in enumerate(self.columns, 1):
                max_length = max(
                    [len(str(ws.cell(row=row, column=col_num).value)) for row in range(1, ws.max_row + 1)]
                )
                ws.column_dimensions[get_column_letter(col_num)].width = max_length + 2

            wb.save(file_path)
            messagebox.showinfo("완료", f"엑셀 파일로 저장되었습니다.\n{file_path}")

        except Exception as e:
            messagebox.showerror("오류", f"엑셀 저장 중 오류 발생:\n{e}")